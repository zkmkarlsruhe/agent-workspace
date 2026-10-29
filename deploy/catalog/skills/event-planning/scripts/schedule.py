# SPDX-License-Identifier: MIT

"""Event schedule worked back from the date, as xlsx (sheet texts in German).

Usage: python3 schedule.py <date DD.MM.YYYY> <target.xlsx> [title]
Due dates on a Saturday or Sunday move to the Friday before.
"""
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

# (days before the date, area, task); negative = after the date. German: written into the sheet.
TASKS = [
    (84, "Organisation", "Ziel, Format, Zielgruppe und Budgetrahmen festlegen"),
    (84, "Organisation", "Raum und Termin verbindlich reservieren"),
    (70, "Technik", "Technikbedarf (Ton, Licht, Projektion, Strom, Netz) melden"),
    (70, "Finanzen", "Budget und Kostenstelle abstimmen"),
    (63, "Programm", "Mitwirkende anfragen, Honorare und Verträge klären"),
    (56, "Recht", "Genehmigungen und Auflagen prüfen (Sicherheit, GEMA, Bild- und Filmrechte)"),
    (56, "Barrierefreiheit", "Zugang, Höranlage, Gebärdensprache, Texte in Einfacher Sprache prüfen"),
    (42, "Kommunikation", "Ankündigung und Pressetext an die Kommunikation geben"),
    (42, "Kommunikation", "Termin für Website und Veranstaltungskalender melden"),
    (35, "Gäste", "Einladung versenden"),
    (28, "Catering", "Angebote einholen und beauftragen"),
    (21, "Reisen", "Anreise und Unterkunft der Mitwirkenden buchen"),
    (14, "Haus", "Ablaufplan an Technik, Pforte, Aufsicht und Reinigung geben"),
    (14, "Gäste", "Rückmeldungen auswerten, Gästeliste führen"),
    (7, "Programm", "Ablauf, Moderation und Begrüßung festlegen"),
    (7, "Kommunikation", "Erinnerung an Gäste und Presse senden"),
    (2, "Technik", "Technikprobe vor Ort"),
    (1, "Organisation", "Beschilderung, Namensschilder und Material bereitlegen"),
    (0, "Organisation", "Veranstaltung"),
    (-2, "Finanzen", "Rechnungen sammeln, Honorare zur Zahlung geben"),
    (-7, "Kommunikation", "Fotos, Nachbericht, Dank an Mitwirkende"),
    (-14, "Organisation", "Auswertung: Besuchszahl, Budget Soll und Ist, Lehren"),
]
WEEKDAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]
HEADER = ["Nr.", "Fällig am", "Wochentag", "Tage vorher", "Bereich", "Aufgabe", "Verantwortlich", "Status", "Notiz"]


def due(day: date, days: int) -> date:
    d = day - timedelta(days=days)
    if days != 0 and d.weekday() >= 5:
        d -= timedelta(days=d.weekday() - 4)
    return d


def rows(day: date):
    for nr, (days, area, task) in enumerate(sorted(TASKS, key=lambda a: -a[0]), 1):
        d = due(day, days)
        yield [nr, d, WEEKDAYS[d.weekday()], days, area, task, "", "offen", ""]


def write(day: date, target: Path, title: str) -> int:
    if target.exists():
        raise FileExistsError(f"{target} exists; get a free name with next_name.py")
    wb = Workbook()
    ws = wb.active
    ws.title = "Zeitplan"
    ws.append([f"{title} am {day:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(HEADER)
    for c in ws[2]:
        c.font = Font(bold=True)
    n = 0
    for row in rows(day):
        ws.append(row)
        ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"
        n += 1
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:I{ws.max_row}"
    for col, width in zip("ABCDEFGHI", [5, 12, 12, 11, 16, 70, 20, 10, 30]):
        ws.column_dimensions[col].width = width
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return n


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    day = datetime.strptime(sys.argv[1], "%d.%m.%Y").date()
    title = sys.argv[3] if len(sys.argv) > 3 else "Veranstaltung"
    target = Path(sys.argv[2])
    n = write(day, target, title)
    past = [r for r in rows(day) if r[1] < date.today() and r[3] > 0]
    print(f"{target}: {n} tasks")
    if past:
        print(f"Already past: {len(past)} tasks due before today, e.g. {past[0][5]} ({past[0][1]:%d.%m.%Y})")


if __name__ == "__main__":
    main()
