# SPDX-License-Identifier: MIT

"""Test for schedule.py. Usage: python3 -I -B test_schedule.py (in a venv with openpyxl)."""
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[2] / "skills" / "event-planning" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from schedule import TASKS, due, write

from openpyxl import load_workbook

# Saturday 2026-12-12: 14 days before is Saturday 28.11. -> Friday 27.11.; the date itself stays Saturday
t = date(2026, 12, 12)
assert due(t, 0) == t
assert due(t, 14) == date(2026, 11, 27)
assert due(t, 1) == date(2026, 12, 11)
assert due(t, -2) == date(2026, 12, 14)  # Monday stays
assert all(due(t, a[0]).weekday() < 5 for a in TASKS if a[0] != 0)

d = Path(tempfile.mkdtemp())
target = d / "results" / "schedule.xlsx"
assert write(t, target, "Lesung") == len(TASKS)
ws = load_workbook(target).active
assert ws["A1"].value == "Lesung am 12.12.2026"
assert ws["F2"].value == "Aufgabe"
values = [ws.cell(r, 2).value for r in range(3, ws.max_row + 1)]
assert values == sorted(values), "sorted by date"
assert ws.cell(3, 2).number_format == "DD.MM.YYYY"
assert isinstance(ws.cell(3, 2).value, datetime)

try:
    write(t, target, "x")
    raise AssertionError("must not overwrite")
except FileExistsError:
    pass

r = subprocess.run([sys.executable, SCRIPTS / "schedule.py", "12.12.2026", str(d / "b.xlsx"), "Lesung"],
                   capture_output=True, text=True)
assert r.returncode == 0 and "22 tasks" in r.stdout, r
print("schedule ok")
