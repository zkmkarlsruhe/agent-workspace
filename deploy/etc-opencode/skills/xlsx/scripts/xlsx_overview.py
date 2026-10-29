#!/usr/bin/env python3
# SPDX-License-Identifier: MIT

"""Overview of an xlsx file: sheets, size, first rows, formulas, merged cells.

Usage: python3 xlsx_overview.py inbox/file.xlsx
"""
import sys

from openpyxl import load_workbook


def overview(path: str, rows: int = 3) -> str:
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    out = []
    for ws in formulas.worksheets:
        wv = values[ws.title]
        formula_count = 0
        no_value = []
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    formula_count += 1
                    if wv[c.coordinate].value is None:
                        no_value.append(c.coordinate)
        out.append(f"## Blatt '{ws.title}': {ws.dimensions}, {ws.max_row} Zeilen x {ws.max_column} Spalten")
        out.append(f"Formeln: {formula_count}, davon ohne gespeicherten Wert: {len(no_value)}"
                   + (f" (z. B. {', '.join(no_value[:5])})" if no_value else ""))
        if ws.merged_cells.ranges:
            out.append("Verbundene Zellen: " + ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:10]))
        for row in wv.iter_rows(max_row=rows, values_only=True):
            out.append("  " + " | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: xlsx_overview.py <file.xlsx>")
    print(overview(sys.argv[1]))
