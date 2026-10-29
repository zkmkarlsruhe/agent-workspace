# SPDX-License-Identifier: MIT

"""Workbook summaries include uncached formulas and empty sheets."""
import importlib.util
from pathlib import Path
import runpy
import tempfile
import unittest

SCRIPT = Path(__file__).resolve().parents[1] / "skills/xlsx/scripts/xlsx_overview.py"


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "requires openpyxl")
class XLSXOverviewTest(unittest.TestCase):
    def test_populated_and_empty_sheets(self):
        from openpyxl import Workbook

        overview = runpy.run_path(str(SCRIPT))["overview"]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            book = Workbook()
            self.addCleanup(book.close)
            sheet = book.active
            sheet.title = "Data"
            sheet.append(["Month", "Total"])
            sheet.append(["Jan", 10])
            sheet.append(["Sum", "=SUM(B2:B2)"])
            sheet.merge_cells("C1:D1")
            book.create_sheet("Empty")
            book.save(path)
            result = overview(str(path))
            self.assertIn("Blatt 'Data': A1:D3, 3 Zeilen x 4 Spalten", result)
            self.assertIn("Jan | 10", result)
            self.assertIn("Formeln: 1, davon ohne gespeicherten Wert: 1 (z. B. B3)", result)
            self.assertIn("Verbundene Zellen: C1:D1", result)
            empty = result.split("## Blatt 'Empty':", 1)[1]
            self.assertIn("A1:A1, 1 Zeilen x 1 Spalten", empty)
            self.assertIn("Formeln: 0, davon ohne gespeicherten Wert: 0", empty)


if __name__ == "__main__":
    unittest.main()
